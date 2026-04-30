"""Parser nodes: `select_parser` + per-format extractors.

Each parser:

1. reads `state.local_path` (downloaded by `select_parser`),
2. extracts text blocks (with structural-tag markers for tables /
   speaker notes / image references),
3. uploads any embedded images to S3 and records `ImageItem`s.
"""

from __future__ import annotations

import hashlib
import os
import tempfile
from typing import Any

import fitz  # PyMuPDF
import pandas as pd
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.graph.nodes.status import node
from app.graph.state import AgentState, ImageItem, TextBlock
from app.services.s3_service import S3Client

_s3 = S3Client()

# Map document type → file extension on disk.
_EXT_BY_TYPE = {"PDF": ".pdf", "DOCX": ".docx", "PPT": ".pptx", "EXCEL": ".xlsx"}

# Map document type → registered node name (used by `route_after_select_parser`).
PARSER_BY_DOC_TYPE = {
    "PDF": "parse_pdf",
    "DOCX": "parse_docx",
    "PPT": "parse_pptx",
    "EXCEL": "parse_excel",
}


# ─────────────────────────────────────────────────────────────────────────────
#  STRUCTURAL-TAG HELPERS  (must match patterns the chunker protects)
# ─────────────────────────────────────────────────────────────────────────────


def _img_placeholder(page: int, image_index: int, s3_url: str) -> str:
    return (
        f'<IMAGE_REFERENCE src="{s3_url}" page="{page}" index="{image_index}">'
        f"</IMAGE_REFERENCE>"
    )


def _table_block(df: pd.DataFrame) -> str:
    return f"<TABLE_REFERENCE>\n{df.to_csv(index=False).strip()}\n</TABLE_REFERENCE>"


def _notes_block(notes: str) -> str:
    return f"<SPEAKER_NOTES>{notes.strip()}</SPEAKER_NOTES>"


def _sha256_of_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# ─────────────────────────────────────────────────────────────────────────────
#  SELECT PARSER  — download + verify the file
# ─────────────────────────────────────────────────────────────────────────────


@node("select_parser")
def select_parser_node(state: AgentState) -> dict[str, Any]:
    """Validate `doc_type`, download binary from S3, and verify SHA-256.

    The actual parser dispatch happens in `route_after_select_parser`.
    """
    meta = state["doc_metadata"]
    if meta.doc_type not in _EXT_BY_TYPE:
        raise ValueError(f"Unsupported doc_type: {meta.doc_type!r}")

    suffix = _EXT_BY_TYPE[meta.doc_type]
    fd, dest_path = tempfile.mkstemp(prefix=f"ingest-{meta.doc_id}-", suffix=suffix)
    os.close(fd)
    _s3.download_to_path(meta.s3_url, dest_path)

    actual = _sha256_of_file(dest_path)
    if actual.lower() != meta.doc_hash.lower():
        raise ValueError(
            f"doc_hash mismatch: expected={meta.doc_hash} actual={actual} for {meta.s3_url}"
        )

    return {"local_path": dest_path}


# ─────────────────────────────────────────────────────────────────────────────
#  PARSERS
# ─────────────────────────────────────────────────────────────────────────────


@node("parse_pdf")
def parse_pdf_node(state: AgentState) -> dict[str, Any]:
    """Extract text + images from a PDF using PyMuPDF; upload images to S3."""
    meta = state["doc_metadata"]
    path = state["local_path"]

    text_blocks: list[TextBlock] = []
    image_items: list[ImageItem] = []

    pdf = fitz.open(path)
    try:
        for page_no, page in enumerate(pdf, start=1):
            page_text = (page.get_text("text") or "").strip()
            placeholders: list[str] = []
            for img_idx, img in enumerate(page.get_images(full=True), start=1):
                xref = img[0]
                base = pdf.extract_image(xref)
                ext = base.get("ext", "png")
                s3_url = _s3.upload_image(
                    meta.knowledge_base_type, meta.doc_id, page_no, img_idx, base["image"], ext=ext
                )
                placeholders.append(_img_placeholder(page_no, img_idx, s3_url))
                image_items.append(
                    ImageItem(page=page_no, image_index=img_idx, s3_url=s3_url)
                )
            combined = page_text
            if placeholders:
                combined = (combined + "\n" + "\n".join(placeholders)).strip()
            if combined:
                text_blocks.append(
                    TextBlock(page=page_no, section=f"page-{page_no}", text=combined)
                )
    finally:
        pdf.close()

    return {"extracted_text_blocks": text_blocks, "extracted_image_items": image_items}


@node("parse_docx")
def parse_docx_node(state: AgentState) -> dict[str, Any]:
    """Extract paragraphs (grouped by heading) + embedded images from a DOCX."""
    meta = state["doc_metadata"]
    path = state["local_path"]

    docx = DocxDocument(path)
    text_blocks: list[TextBlock] = []
    image_items: list[ImageItem] = []
    section_lines: list[str] = []
    current_section = "section-1"
    section_page = 1

    for para in docx.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if style.startswith("heading"):
            if section_lines:
                text_blocks.append(
                    TextBlock(
                        page=section_page,
                        section=current_section,
                        text="\n".join(section_lines),
                    )
                )
                section_lines = []
                section_page += 1
            current_section = text
        else:
            section_lines.append(text)

    img_idx = 0
    for rel in docx.part._rels.values():
        if "image" not in rel.reltype:
            continue
        img_idx += 1
        blob = rel.target_part.blob
        ext = (rel.target_part.content_type.split("/")[-1] or "png").lower()
        s3_url = _s3.upload_image(
            meta.knowledge_base_type, meta.doc_id, page=1, image_index=img_idx, data=blob, ext=ext
        )
        section_lines.append(_img_placeholder(1, img_idx, s3_url))
        image_items.append(ImageItem(page=1, image_index=img_idx, s3_url=s3_url))

    if section_lines:
        text_blocks.append(
            TextBlock(page=section_page, section=current_section, text="\n".join(section_lines))
        )

    return {"extracted_text_blocks": text_blocks, "extracted_image_items": image_items}


@node("parse_pptx")
def parse_pptx_node(state: AgentState) -> dict[str, Any]:
    """Extract slides, pictures, and SPEAKER_NOTES from a PPTX."""
    meta = state["doc_metadata"]
    path = state["local_path"]

    prs = Presentation(path)
    text_blocks: list[TextBlock] = []
    image_items: list[ImageItem] = []

    for slide_no, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        img_idx = 0
        for shape in slide.shapes:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    lines.append(txt)
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_idx += 1
                image = shape.image
                ext = (image.ext or "png").lower()
                s3_url = _s3.upload_image(
                    meta.knowledge_base_type,
                    meta.doc_id,
                    slide_no,
                    img_idx,
                    image.blob,
                    ext=ext,
                )
                lines.append(_img_placeholder(slide_no, img_idx, s3_url))
                image_items.append(
                    ImageItem(page=slide_no, image_index=img_idx, s3_url=s3_url)
                )
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(_notes_block(notes))
        if lines:
            text_blocks.append(
                TextBlock(page=slide_no, section=f"slide-{slide_no}", text="\n".join(lines))
            )

    return {"extracted_text_blocks": text_blocks, "extracted_image_items": image_items}


@node("parse_excel")
def parse_excel_node(state: AgentState) -> dict[str, Any]:
    """Extract sheet rows (TABLE_REFERENCE) + embedded images from an XLSX."""
    meta = state["doc_metadata"]
    path = state["local_path"]

    text_blocks: list[TextBlock] = []
    image_items: list[ImageItem] = []
    rows_per_block = 50

    sheets = pd.read_excel(path, sheet_name=None, dtype=object)
    for sheet_index, (sheet_name, df) in enumerate(sheets.items(), start=1):
        df = df.where(pd.notna(df), "")
        for start in range(0, len(df), rows_per_block):
            slice_df = df.iloc[start : start + rows_per_block]
            text_blocks.append(
                TextBlock(
                    page=sheet_index,
                    section=f"{sheet_name}:rows-{start + 1}-{start + len(slice_df)}",
                    text=_table_block(slice_df),
                )
            )

    wb = load_workbook(path, data_only=True)
    try:
        for sheet_index, ws in enumerate(wb.worksheets, start=1):
            for img_idx, image in enumerate(getattr(ws, "_images", []) or [], start=1):
                blob = image._data() if callable(getattr(image, "_data", None)) else None
                if not blob:
                    continue
                s3_url = _s3.upload_image(
                    meta.knowledge_base_type, meta.doc_id, sheet_index, img_idx, blob, ext="png"
                )
                text_blocks.append(
                    TextBlock(
                        page=sheet_index,
                        section=f"{ws.title}:image-{img_idx}",
                        text=_img_placeholder(sheet_index, img_idx, s3_url),
                    )
                )
                image_items.append(
                    ImageItem(page=sheet_index, image_index=img_idx, s3_url=s3_url)
                )
    finally:
        wb.close()

    return {"extracted_text_blocks": text_blocks, "extracted_image_items": image_items}
